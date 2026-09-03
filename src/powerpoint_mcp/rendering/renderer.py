"""PowerPoint presentation rendering pipeline supporting Windows COM and LibreOffice headless fallback."""

from abc import ABC, abstractmethod
import gc
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
from typing import Any, Dict, List, Optional, Union


class BaseRenderer(ABC):
    """Abstract base class for PowerPoint slide and presentation renderers."""

    @property
    @abstractmethod
    def renderer_name(self) -> str:
        """Name of the rendering engine (e.g., 'powerpoint', 'libreoffice')."""
        ...

    @property
    @abstractmethod
    def is_available(self) -> bool:
        """Check if the renderer engine and host dependencies are available."""
        ...

    def get_renderer_info(self) -> Dict[str, Any]:
        """Return diagnostic metadata about the renderer."""
        return {
            "renderer_name": self.renderer_name,
            "is_available": self.is_available,
            "platform": sys.platform,
        }

    @abstractmethod
    def render_slide(
        self,
        presentation_path: Union[str, Path],
        slide_number: int,
        output_path: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> str:
        """Render a single 1-indexed slide from a presentation to a PNG image file.

        Args:
            presentation_path: Path to the .pptx presentation file.
            slide_number: 1-indexed slide number.
            output_path: Target PNG output file path.
            width: Output image width in pixels (default: 1920).
            height: Output image height in pixels (default: 1080).

        Returns:
            Absolute path string to the generated PNG image.
        """
        ...

    @abstractmethod
    def render_presentation(
        self,
        presentation_path: Union[str, Path],
        output_dir: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> List[str]:
        """Render all slides of a presentation to PNG files in output_dir.

        Args:
            presentation_path: Path to the .pptx presentation file.
            output_dir: Directory where slide PNG files will be written.
            width: Output image width in pixels (default: 1920).
            height: Output image height in pixels (default: 1080).

        Returns:
            List of absolute path strings to the generated PNG images in slide order.
        """
        ...

    def render_slides(
        self,
        presentation_path: Union[str, Path],
        slide_numbers: List[int],
        output_dir: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> Dict[int, str]:
        """Render a specific list of slides to PNG files.

        Args:
            presentation_path: Path to presentation.
            slide_numbers: List of 1-indexed slide numbers.
            output_dir: Directory where slide PNG files will be written.
            width: Output image width in pixels.
            height: Output image height in pixels.

        Returns:
            Dictionary mapping slide number to absolute path string.
        """
        results: Dict[int, str] = {}
        for s_num in slide_numbers:
            out_p = Path(output_dir) / f"slide_{s_num}.png"
            results[s_num] = self.render_slide(presentation_path, s_num, out_p, width, height)
        return results


def _render_slide_inner(
    ppt_app: Any, prs_path: Path, slide_number: int, out_path: Path, width: int, height: int
) -> None:
    import pythoncom

    presentations = ppt_app.Presentations
    # Open(FileName, ReadOnly=-1, Untitled=0, WithWindow=0)
    presentation = presentations.Open(str(prs_path), -1, 0, 0)
    try:
        slides = presentation.Slides
        slide_count = slides.Count
        if slide_number < 1 or slide_number > slide_count:
            raise IndexError(
                f"Slide number {slide_number} is out of range. Presentation contains {slide_count} slides."
            )
        slide = slides.Item(slide_number)
        try:
            slide.Export(str(out_path), "PNG", int(width), int(height))
        finally:
            del slide
        del slides
    finally:
        try:
            presentation.Saved = -1
            presentation.Close()
        except Exception:
            pass
        del presentation
        del presentations
        try:
            pythoncom.PumpWaitingMessages()
        except Exception:
            pass
        gc.collect()


def _com_export_slide(
    prs_path: Path, slide_number: int, out_path: Path, width: int, height: int
) -> None:
    from powerpoint_mcp.rendering.com_lifecycle import com_powerpoint_session

    with com_powerpoint_session() as (app, _):
        _render_slide_inner(app, prs_path, slide_number, out_path, width, height)
        del app


def _render_presentation_inner(
    ppt_app: Any, prs_path: Path, out_dir: Path, width: int, height: int
) -> List[str]:
    import pythoncom

    rendered_paths: List[str] = []
    presentations = ppt_app.Presentations
    # Open(FileName, ReadOnly=-1, Untitled=0, WithWindow=0)
    presentation = presentations.Open(str(prs_path), -1, 0, 0)
    try:
        slides = presentation.Slides
        slide_count = slides.Count
        for idx in range(1, slide_count + 1):
            slide_out = out_dir / f"slide_{idx}.png"
            slide = slides.Item(idx)
            try:
                slide.Export(str(slide_out), "PNG", int(width), int(height))
            finally:
                del slide
            if not slide_out.exists() or slide_out.stat().st_size == 0:
                raise RuntimeError(
                    f"PowerPoint COM export failed for slide {idx} at: {slide_out}"
                )
            rendered_paths.append(str(slide_out))
        del slides
    finally:
        try:
            presentation.Saved = -1
            presentation.Close()
        except Exception:
            pass
        del presentation
        del presentations
        try:
            pythoncom.PumpWaitingMessages()
        except Exception:
            pass
        gc.collect()
    return rendered_paths


def _com_export_presentation(
    prs_path: Path, out_dir: Path, width: int, height: int
) -> List[str]:
    from powerpoint_mcp.rendering.com_lifecycle import com_powerpoint_session

    results: List[str] = []
    with com_powerpoint_session() as (app, _):
        results = _render_presentation_inner(app, prs_path, out_dir, width, height)
        del app
    return results


def _render_slides_inner(
    ppt_app: Any,
    prs_path: Path,
    out_dir: Path,
    slide_numbers: List[int],
    width: int,
    height: int,
) -> Dict[int, str]:
    import pythoncom

    rendered_map: Dict[int, str] = {}
    presentations = ppt_app.Presentations
    # Open(FileName, ReadOnly=-1, Untitled=0, WithWindow=0)
    presentation = presentations.Open(str(prs_path), -1, 0, 0)
    try:
        slides = presentation.Slides
        total_count = slides.Count
        for idx in slide_numbers:
            if idx < 1 or idx > total_count:
                raise IndexError(f"Slide number {idx} is out of range (1..{total_count})")
            slide_out = out_dir / f"slide_{idx}.png"
            slide = slides.Item(idx)
            try:
                slide.Export(str(slide_out), "PNG", int(width), int(height))
            finally:
                del slide
            if not slide_out.exists() or slide_out.stat().st_size == 0:
                raise RuntimeError(
                    f"PowerPoint COM export failed for slide {idx} at: {slide_out}"
                )
            rendered_map[idx] = str(slide_out)
        del slides
    finally:
        try:
            presentation.Saved = -1
            presentation.Close()
        except Exception:
            pass
        del presentation
        del presentations
        try:
            pythoncom.PumpWaitingMessages()
        except Exception:
            pass
        gc.collect()
    return rendered_map


def _com_export_slides(
    prs_path: Path,
    out_dir: Path,
    slide_numbers: List[int],
    width: int,
    height: int,
) -> Dict[int, str]:
    from powerpoint_mcp.rendering.com_lifecycle import com_powerpoint_session

    results: Dict[int, str] = {}
    with com_powerpoint_session() as (app, _):
        results = _render_slides_inner(app, prs_path, out_dir, slide_numbers, width, height)
        del app
    return results


class PowerPointRenderer(BaseRenderer):
    """Native Microsoft PowerPoint COM automation renderer on Windows.

    Uses `win32com.client.DispatchEx('PowerPoint.Application')` with strict
    Single-Threaded Apartment (STA) lifecycle management (`pythoncom.CoInitialize()`,
    `try...finally`, `presentation.Close()`, `ppt_app.Quit()`, `pythoncom.CoUninitialize()`,
    `gc.collect()`). Runs invisibly with zero orphaned processes.
    """

    @property
    def renderer_name(self) -> str:
        return "powerpoint"

    @property
    def is_available(self) -> bool:
        if sys.platform != "win32":
            return False
        try:
            import winreg

            key = winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application")
            winreg.CloseKey(key)
            import pythoncom
            import win32com.client  # noqa: F401

            return True
        except Exception:
            return False

    def render_slide(
        self,
        presentation_path: Union[str, Path],
        slide_number: int,
        output_path: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> str:
        if not self.is_available:
            raise RuntimeError(
                "PowerPoint COM renderer is not available on this platform or PowerPoint is not installed."
            )

        prs_path = Path(presentation_path).resolve()
        if not prs_path.exists():
            raise FileNotFoundError(f"Presentation file does not exist: {prs_path}")

        if slide_number < 1:
            raise IndexError(f"Slide number must be >= 1 (1-indexed), got {slide_number}")

        from pptx import Presentation as PptxPresentation

        try:
            deck = PptxPresentation(str(prs_path))
            total_slides = len(deck.slides)
            if slide_number > total_slides:
                raise IndexError(
                    f"Slide number {slide_number} is out of range. Presentation contains {total_slides} slides."
                )
        except IndexError:
            raise
        except Exception:
            pass

        out_path = Path(output_path).resolve()
        out_path.parent.mkdir(parents=True, exist_ok=True)

        _com_export_slide(prs_path, slide_number, out_path, width, height)

        if not out_path.exists() or out_path.stat().st_size == 0:
            raise RuntimeError(f"PowerPoint COM export failed to produce valid PNG at: {out_path}")

        return str(out_path)

    def render_presentation(
        self,
        presentation_path: Union[str, Path],
        output_dir: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> List[str]:
        if not self.is_available:
            raise RuntimeError(
                "PowerPoint COM renderer is not available on this platform or PowerPoint is not installed."
            )

        prs_path = Path(presentation_path).resolve()
        if not prs_path.exists():
            raise FileNotFoundError(f"Presentation file does not exist: {prs_path}")

        out_dir = Path(output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

        return _com_export_presentation(prs_path, out_dir, width, height)

    def render_slides(
        self,
        presentation_path: Union[str, Path],
        slide_numbers: List[int],
        output_dir: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> Dict[int, str]:
        if not self.is_available:
            raise RuntimeError(
                "PowerPoint COM renderer is not available on this platform or PowerPoint is not installed."
            )

        prs_path = Path(presentation_path).resolve()
        if not prs_path.exists():
            raise FileNotFoundError(f"Presentation file does not exist: {prs_path}")

        out_dir = Path(output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

        return _com_export_slides(prs_path, out_dir, slide_numbers, width, height)


class LibreOfficeRenderer(BaseRenderer):
    """Headless LibreOffice (`soffice`) fallback renderer.

    Converts presentations to PDF headlessly using LibreOffice, then rasterizes
    the PDF slides into high-resolution PNG images.
    """

    def __init__(self, executable_path: Optional[str] = None):
        self._custom_executable = executable_path

    @property
    def renderer_name(self) -> str:
        return "libreoffice"

    def _find_executable(self) -> Optional[str]:
        if self._custom_executable and Path(self._custom_executable).is_file():
            return str(Path(self._custom_executable).resolve())

        which_soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if which_soffice:
            return which_soffice

        candidates = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/usr/bin/soffice",
            "/usr/bin/libreoffice",
            "/usr/local/bin/soffice",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ]
        for candidate in candidates:
            if Path(candidate).is_file():
                return str(Path(candidate).resolve())
        return None

    @property
    def is_available(self) -> bool:
        return self._find_executable() is not None

    def get_renderer_info(self) -> Dict[str, Any]:
        info = super().get_renderer_info()
        info["executable_path"] = self._find_executable()
        return info

    def _convert_to_pdf(self, prs_path: Path, temp_dir: Path) -> Path:
        soffice_bin = self._find_executable()
        if not soffice_bin:
            raise RuntimeError("LibreOffice executable (soffice) not found on system.")

        cmd = [
            soffice_bin,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(temp_dir),
            str(prs_path),
        ]
        try:
            proc = subprocess.run(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=60,
                check=False,
            )
            if proc.returncode != 0:
                raise RuntimeError(
                    f"LibreOffice conversion failed (exit code {proc.returncode}): {proc.stderr}"
                )
        except subprocess.TimeoutExpired:
            raise TimeoutError("LibreOffice conversion timed out after 60 seconds.")

        pdf_path = temp_dir / f"{prs_path.stem}.pdf"
        if not pdf_path.exists():
            raise FileNotFoundError(f"Expected converted PDF not found at: {pdf_path}")
        return pdf_path

    def _rasterize_pdf_page(
        self, pdf_path: Path, page_index: int, output_path: Path, width: int, height: int
    ) -> Path:
        """Rasterize a single PDF page (0-indexed) to PNG."""
        # Strategy 1: pypdfium2
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(pdf_path))
            if page_index < 0 or page_index >= len(pdf):
                raise IndexError(f"Page index {page_index} out of bounds for PDF with {len(pdf)} pages.")
            page = pdf[page_index]
            page_w, page_h = page.get_size()
            scale = max(width / page_w, height / page_h) if page_w > 0 and page_h > 0 else 2.0
            bitmap = page.render(scale=scale)
            pil_image = bitmap.to_pil()
            pil_image = pil_image.resize((width, height))
            pil_image.save(str(output_path), format="PNG")
            return output_path
        except ImportError:
            pass

        # Strategy 2: fitz (PyMuPDF)
        try:
            import fitz

            doc = fitz.open(str(pdf_path))
            if page_index < 0 or page_index >= len(doc):
                raise IndexError(f"Page index {page_index} out of bounds for PDF with {len(doc)} pages.")
            page = doc[page_index]
            zoom_x = width / page.rect.width if page.rect.width > 0 else 2.0
            zoom_y = height / page.rect.height if page.rect.height > 0 else 2.0
            mat = fitz.Matrix(zoom_x, zoom_y)
            pix = page.get_pixmap(matrix=mat)
            pix.save(str(output_path))
            return output_path
        except ImportError:
            pass

        # Strategy 3: pdf2image
        try:
            from pdf2image import convert_from_path

            images = convert_from_path(
                str(pdf_path),
                first_page=page_index + 1,
                last_page=page_index + 1,
                size=(width, height),
            )
            if images:
                images[0].save(str(output_path), format="PNG")
                return output_path
        except ImportError:
            pass

        raise RuntimeError(
            "LibreOffice converted the presentation to PDF, but no supported PDF rasterizer "
            "(pypdfium2, pymupdf/fitz, or pdf2image) is installed in the Python environment."
        )

    def render_slide(
        self,
        presentation_path: Union[str, Path],
        slide_number: int,
        output_path: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> str:
        if not self.is_available:
            raise RuntimeError("LibreOffice executable (soffice) not found on system.")

        prs_path = Path(presentation_path).resolve()
        if not prs_path.exists():
            raise FileNotFoundError(f"Presentation file does not exist: {prs_path}")

        if slide_number < 1:
            raise IndexError(f"Slide number must be >= 1 (1-indexed), got {slide_number}")

        out_path = Path(output_path).resolve()
        out_path.parent.mkdir(parents=True, exist_ok=True)

        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            pdf_path = self._convert_to_pdf(prs_path, Path(tmp_dir))
            self._rasterize_pdf_page(pdf_path, slide_number - 1, out_path, width, height)

        if not out_path.exists() or out_path.stat().st_size == 0:
            raise RuntimeError(f"LibreOffice rendering failed to produce image at: {out_path}")

        return str(out_path)

    def render_presentation(
        self,
        presentation_path: Union[str, Path],
        output_dir: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> List[str]:
        if not self.is_available:
            raise RuntimeError("LibreOffice executable (soffice) not found on system.")

        prs_path = Path(presentation_path).resolve()
        if not prs_path.exists():
            raise FileNotFoundError(f"Presentation file does not exist: {prs_path}")

        out_dir = Path(output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

        import tempfile

        rendered_paths: List[str] = []
        with tempfile.TemporaryDirectory() as tmp_dir:
            pdf_path = self._convert_to_pdf(prs_path, Path(tmp_dir))

            page_count = 0
            try:
                import pypdfium2 as pdfium

                pdf = pdfium.PdfDocument(str(pdf_path))
                page_count = len(pdf)
            except ImportError:
                try:
                    import fitz

                    doc = fitz.open(str(pdf_path))
                    page_count = len(doc)
                except ImportError:
                    try:
                        from pdf2image import pdfinfo_from_path

                        info = pdfinfo_from_path(str(pdf_path))
                        page_count = info.get("Pages", 1)
                    except Exception:
                        page_count = 1

            for idx in range(1, page_count + 1):
                slide_out = out_dir / f"slide_{idx}.png"
                self._rasterize_pdf_page(pdf_path, idx - 1, slide_out, width, height)
                rendered_paths.append(str(slide_out))

        return rendered_paths


class NullRenderer(BaseRenderer):
    """Fallback renderer when no presentation rendering engine is installed on host."""

    @property
    def renderer_name(self) -> str:
        return "none"

    @property
    def is_available(self) -> bool:
        return False

    def render_slide(
        self,
        presentation_path: Union[str, Path],
        slide_number: int,
        output_path: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> str:
        raise RuntimeError(
            "No presentation renderer is available on this system. "
            "Please install Microsoft PowerPoint or LibreOffice."
        )

    def render_presentation(
        self,
        presentation_path: Union[str, Path],
        output_dir: Union[str, Path],
        width: int = 1920,
        height: int = 1080,
    ) -> List[str]:
        raise RuntimeError(
            "No presentation renderer is available on this system. "
            "Please install Microsoft PowerPoint or LibreOffice."
        )


def get_available_renderer(preferred: str = "auto") -> BaseRenderer:
    """Detect and return an appropriate BaseRenderer instance.

    Args:
        preferred: One of 'auto', 'powerpoint', 'libreoffice', or 'none'.
            If 'auto', checks PPT_RENDERER environment variable, then attempts
            PowerPoint COM automation on Windows, falling back to LibreOffice headless.

    Returns:
        A BaseRenderer instance (which has `.renderer_name` and `.is_available`).
    """
    mode = (preferred or "auto").strip().lower()

    if mode == "powerpoint":
        return PowerPointRenderer()
    elif mode == "libreoffice":
        return LibreOfficeRenderer()
    elif mode == "none":
        return NullRenderer()

    # Mode is 'auto' - inspect environment variable
    env_pref = os.environ.get("PPT_RENDERER", "auto").strip().lower()
    if env_pref == "powerpoint":
        return PowerPointRenderer()
    elif env_pref == "libreoffice":
        return LibreOfficeRenderer()
    elif env_pref == "none":
        return NullRenderer()

    # Priority 1: PowerPoint COM (if available)
    ppt = PowerPointRenderer()
    if ppt.is_available:
        return ppt

    # Priority 2: LibreOffice (if available)
    lo = LibreOfficeRenderer()
    if lo.is_available:
        return lo

    # Fallback: NullRenderer
    return NullRenderer()
